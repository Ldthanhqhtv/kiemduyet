"""
Unit test cho core/docx_exporter.py
===================================

Chay: & ".venv\\Scripts\\python.exe" -m unittest core.test_docx_exporter -v

Test tao file .docx that, sau do mo lai bang python-docx de kiem tra noi dung,
cau truc bang, ngat trang, header/footer va font chu.
"""

from __future__ import annotations

import unittest
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

from core.docx_exporter import (
    EDITED_TEXT_HEADING,
    FACT_TABLE_HEADERS,
    HIGHLIGHT_DRAFT_HEADING,
    REPORT_TITLE,
    SPELLING_TABLE_HEADERS,
    SUGGESTED_TITLES_SECTION_HEADING,
    SKILL_NOTES_SECTION_HEADING,
    build_report_document,
    coerce_review_result,
    export_review_to_bytes,
    export_review_to_docx,
    extract_article_title,
)
from core.reviewer import ReviewResult

ORIGINAL_TEXT = """Cháy lớn tại khu công nghiệp Tân Bình

Vào lúc 14 giờ ngày 30/2/2024, một vụ cháy lớn đã xảy ra tại khu công nghiệp Tân Bình, TP.HCM.
Theo thông tin ban đầu, không có thiệt hại về người.
"""

REVIEW_RESULT = {
    "spelling_errors": [
        {
            "loi_sai": "30/2/2024",
            "vi_tri": "câu thứ 1, đoạn 2",
            "goi_y_sua": "28/2/2024",
            "ly_do": "Tháng 2 năm 2024 chỉ có 29 ngày (năm nhuận).",
        },
        {
            "loi_sai": "khu công nghiệp Tân Bình",
            "vi_tri": "câu thứ 1",
            "goi_y_sua": "Khu công nghiệp Tân Bình",
            "ly_do": "Tên địa danh riêng phải viết hoa.",
        },
    ],
    "fact_checks": [
        {
            "thong_tin_can_check": "Vụ cháy tại khu công nghiệp Tân Bình tháng 2/2024",
            "ket_qua_xac_minh": "ĐÚNG - có vụ cháy được báo chí đưa tin.",
            "trich_dan_nguon": "https://vnexpress.net/example-1",
            "do_tin_cay": "Cao",
        }
    ],
    "edited_text": "Cháy lớn tại Khu công nghiệp Tân Bình\n\nVào lúc 14 giờ ngày 28/2/2024, một vụ cháy lớn đã xảy ra tại Khu công nghiệp Tân Bình, TP.HCM.",
    "grounding_sources": [
        {
            "title": "VnExpress - Vụ cháy khu công nghiệp",
            "uri": "https://vnexpress.net/example-1",
            "domain": "vnexpress.net",
        }
    ],
    "search_queries": ["cháy khu công nghiệp Tân Bình 2024"],
    "model": "gemini-2.5-flash",
}


def _all_text(doc: Document) -> str:
    """Gom toan bo van ban trong tai lieu (doan van + o bang)."""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


class TestCoerceAndHelpers(unittest.TestCase):
    def test_coerce_dict(self) -> None:
        result = coerce_review_result(REVIEW_RESULT)
        self.assertEqual(len(result["spelling_errors"]), 2)
        self.assertEqual(len(result["fact_checks"]), 1)

    def test_coerce_reviewresult_dataclass(self) -> None:
        obj = ReviewResult(
            spelling_errors=[{"loi_sai": "a"}],
            fact_checks=[],
            edited_text="van ban",
            model="gemini-2.5-flash",
        )
        result = coerce_review_result(obj)
        self.assertEqual(result["edited_text"], "van ban")
        self.assertEqual(result["spelling_errors"][0]["loi_sai"], "a")

    def test_coerce_none(self) -> None:
        result = coerce_review_result(None)
        self.assertEqual(result["spelling_errors"], [])
        self.assertEqual(result["edited_text"], "")

    def test_extract_article_title_uses_first_line(self) -> None:
        self.assertEqual(
            extract_article_title(ORIGINAL_TEXT), "Cháy lớn tại khu công nghiệp Tân Bình"
        )

    def test_extract_article_title_empty(self) -> None:
        self.assertEqual(extract_article_title(""), "(Không có tiêu đề)")


class TestExportDocument(unittest.TestCase):
    def setUp(self) -> None:
        self.out_dir = Path(__file__).parent / "_tmp_docx"
        self.out_dir.mkdir(exist_ok=True)
        self.path = self.out_dir / "report.docx"
        self.fixed_dt = datetime(2026, 10, 17, 9, 30)

    def tearDown(self) -> None:
        if self.path.exists():
            self.path.unlink()
        if self.out_dir.exists() and not any(self.out_dir.iterdir()):
            self.out_dir.rmdir()

    def _export(self) -> Document:
        export_review_to_docx(
            ORIGINAL_TEXT,
            REVIEW_RESULT,
            self.path,
            newspaper_name="BÁO ĐIỆN TỬ THỬ NGHIỆM",
            editor_name="Ban Biên Tập",
            export_datetime=self.fixed_dt,
        )
        self.assertTrue(self.path.is_file())
        return Document(str(self.path))

    def test_report_title_and_datetime_present(self) -> None:
        doc = self._export()
        text = _all_text(doc)
        self.assertIn(REPORT_TITLE, text)
        self.assertIn("BÁO ĐIỆN TỬ THỬ NGHIỆM", text)
        self.assertIn("Ngày 17 tháng 10 năm 2026 - 09:30", text)

    def test_two_data_tables_with_correct_headers(self) -> None:
        doc = self._export()
        self.assertGreaterEqual(len(doc.tables), 2)

        spelling_table = doc.tables[0]
        spelling_headers = [cell.text for cell in spelling_table.rows[0].cells]
        self.assertEqual(spelling_headers, list(SPELLING_TABLE_HEADERS))

        fact_table = doc.tables[1]
        fact_headers = [cell.text for cell in fact_table.rows[0].cells]
        self.assertEqual(fact_headers, list(FACT_TABLE_HEADERS))

    def test_spelling_table_contains_errors(self) -> None:
        doc = self._export()
        table = doc.tables[0]
        # 1 dong tieu de + 2 dong du lieu
        self.assertEqual(len(table.rows), 3)
        row_texts = [cell.text for cell in table.rows[1].cells]
        self.assertEqual(row_texts[0], "1")
        self.assertEqual(row_texts[1], "30/2/2024")
        self.assertEqual(row_texts[3], "28/2/2024")

    def test_fact_table_contains_check(self) -> None:
        doc = self._export()
        table = doc.tables[1]
        self.assertEqual(len(table.rows), 2)
        row_texts = [cell.text for cell in table.rows[1].cells]
        self.assertIn("https://vnexpress.net/example-1", row_texts[-1])
        self.assertEqual(row_texts[3], "Cao")

    def test_page_break_present(self) -> None:
        self._export()
        with zipfile.ZipFile(self.path) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn('w:type="page"', document_xml)

    def test_edited_text_section_present(self) -> None:
        doc = self._export()
        text = _all_text(doc)
        self.assertIn(EDITED_TEXT_HEADING, text)
        self.assertIn("Vào lúc 14 giờ ngày 28/2/2024", text)

    def test_grounding_appendix_present(self) -> None:
        doc = self._export()
        text = _all_text(doc)
        self.assertIn("Nguồn tham chiếu tra cứu (Google Search)", text)

    def test_body_font_is_times_new_roman(self) -> None:
        doc = self._export()
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if run.text.strip():
                    self.assertEqual(run.font.name, "Times New Roman")

    def test_header_and_footer(self) -> None:
        doc = self._export()
        header_text = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
        self.assertIn("BÁO ĐIỆN TỬ THỬ NGHIỆM", header_text)
        # Footer chua field PAGE -> kiem tra trong XML.
        with zipfile.ZipFile(self.path) as archive:
            footer_xml = archive.read("word/footer1.xml").decode("utf-8")
        self.assertIn("PAGE", footer_xml)


class TestEmptyAndBytes(unittest.TestCase):
    def test_empty_review_renders_placeholder_rows(self) -> None:
        doc = build_report_document(
            "Bai bao khong co loi.",
            {"spelling_errors": [], "fact_checks": [], "edited_text": ""},
            export_datetime=datetime(2026, 1, 1, 0, 0),
        )
        text = _all_text(doc)
        self.assertIn("Không phát hiện lỗi chính tả.", text)
        self.assertIn("Không có thông tin nào cần kiểm chứng.", text)
        # edited_text rong -> dung fallback original_text.
        self.assertIn("Bai bao khong co loi.", text)

    def test_export_bytes_returns_valid_docx(self) -> None:
        data = export_review_to_bytes(ORIGINAL_TEXT, REVIEW_RESULT)
        self.assertIsInstance(data, bytes)
        self.assertTrue(data.startswith(b"PK"))  # chữ ký zip của file docx
        self.assertGreater(len(data), 1000)


class TestHighlightedDraft(unittest.TestCase):
    """Phan "BẢN THẢO GỐC ĐÁNH DẤU LỖI NGHI NGỜ" (TASK_05)."""

    def _export(self) -> Document:
        return build_report_document(
            ORIGINAL_TEXT,
            REVIEW_RESULT,
            newspaper_name="BÁO ĐIỆN TỬ THỬ NGHIỆM",
            export_datetime=datetime(2026, 10, 18, 10, 0),
        )

    def test_highlighted_draft_heading_present(self) -> None:
        text = _all_text(self._export())
        self.assertIn(HIGHLIGHT_DRAFT_HEADING, text)

    def test_draft_summary_counts_marked_points(self) -> None:
        text = _all_text(self._export())
        # 3 diem: "khu công nghiệp Tân Bình" xuat hien 2 lan (tieu de + doan 2)
        # va "30/2/2024" xuat hien 1 lan; fact-check "ĐÚNG" khong duoc boi do.
        self.assertIn("Đã đánh dấu 3 điểm nghi ngờ sai", text)

    def test_error_runs_are_red_bold_underlined(self) -> None:
        doc = self._export()
        red_runs = []
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if run.font.color and run.font.color.rgb is not None:
                    if str(run.font.color.rgb) == "DC2626":
                        red_runs.append(run)
        self.assertEqual(len(red_runs), 3)
        for run in red_runs:
            self.assertTrue(run.bold)
            self.assertTrue(run.underline)
            self.assertEqual(run.font.name, "Times New Roman")

    def test_original_text_preserved_in_draft(self) -> None:
        text = _all_text(self._export())
        self.assertIn("Vào lúc 14 giờ ngày 30/2/2024", text)
        self.assertIn("khu công nghiệp Tân Bình, TP.HCM.", text)

    def test_option_disables_draft_section(self) -> None:
        doc = build_report_document(
            ORIGINAL_TEXT,
            REVIEW_RESULT,
            include_highlighted_draft=False,
            export_datetime=datetime(2026, 10, 18, 10, 0),
        )
        self.assertNotIn(HIGHLIGHT_DRAFT_HEADING, _all_text(doc))

    def test_draft_keeps_line_breaks(self) -> None:
        """Xuong dong trong bai goc phai duoc giu (moi dong -> mot doan Word)."""
        doc = self._export()
        # Chi dem doan trong PHAN BAN THAO GOC (dong goc co ngay 30/2/2024).
        draft_paragraphs = [
            p.text
            for p in doc.paragraphs
            if p.text.strip() and "30/2/2024" in p.text
        ]
        self.assertEqual(len(draft_paragraphs), 1)
        self.assertIn("TP.HCM.", draft_paragraphs[0])
        # Dong ke sau (Thong tin ban đầu) van la doan rieng — khong bi noi.
        self.assertTrue(
            any("không có thiệt hại về người" in p.text for p in doc.paragraphs)
        )

    def test_empty_original_renders_placeholder(self) -> None:
        doc = build_report_document(
            "",
            {"spelling_errors": [], "fact_checks": [], "edited_text": ""},
            export_datetime=datetime(2026, 1, 1, 0, 0),
        )
        self.assertIn("(Không có nội dung văn bản gốc.)", _all_text(doc))

    def test_export_includes_suggested_titles_and_skill_notes(self) -> None:
        result = {
            "spelling_errors": [],
            "fact_checks": [],
            "suggested_titles": ["Tiêu đề gợi ý 1 chuẩn SEO", "Tiêu đề gợi ý 2"],
            "skill_notes": "Sapo mở đầu bài báo hấp dẫn về công nghiệp Tân Bình.",
            "edited_text": "Bản thảo đã sửa.",
        }
        doc = build_report_document(
            ORIGINAL_TEXT,
            result,
            export_datetime=datetime(2026, 1, 1, 0, 0),
        )
        text = _all_text(doc)
        self.assertIn(SUGGESTED_TITLES_SECTION_HEADING, text)
        self.assertIn("Tiêu đề gợi ý 1 chuẩn SEO", text)
        self.assertIn(SKILL_NOTES_SECTION_HEADING, text)
        self.assertIn("Sapo mở đầu bài báo hấp dẫn", text)
        self.assertIn("Gợi ý tiêu đề: 2", text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
